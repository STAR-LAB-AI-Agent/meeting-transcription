"""Insert real screenshots into the filled internship manual (new file)."""

from __future__ import annotations

from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DOCS = PROJECT_ROOT / "docs"
IMG_DIR = DOCS / "screenshots" / "final_manual"
SRC = DOCS / "网安学院本科实习手册_AI会议语音转写与检索_已填写.docx"
DST = DOCS / "网安学院本科实习手册_AI会议语音转写与检索_最终版.docx"

# (anchor section prefix, image filename, caption)
PLAN = [
    ("三、技术方案与总体架构", "fig1_gui_home.png", "图1 AI会议语音转写与检索智能体主界面"),
    ("五、会议转写和时间戳设计", "fig2_transcript.png", "图2 FunASR会议音频带时间戳转写结果"),
    ("六、关键词检索和时间定位", "fig3_search.png", "图3 “预算”关键词检索结果"),
    ("六、关键词检索和时间定位", "fig4_locate.png", "图4 00:30时间位置查询结果"),
    ("七、自动会议摘要", "fig5_summary.png", "图5 自动会议摘要生成结果"),
    ("八、Agent Router 与 Skill", "fig6_skill.png", "图6 nanobot中meeting-asr Skill加载情况"),
    ("九、nanobot Runtime 集成", "fig7_runtime.png", "图7 nanobot Runtime调用meeting-asr Script/CLI"),
    ("十三、测试与问题解决", "fig8_pytest.png", "图8 pytest自动测试结果"),
]


def _set_run_font(run, size: float = 10.5, font: str = "宋体") -> None:
    run.font.size = docx.shared.Pt(size)
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(qn("w:eastAsia"), font)


def _insert_image(cell, anchor: Paragraph, image: Path, caption: str) -> Paragraph:
    img_el = OxmlElement("w:p")
    anchor._element.addnext(img_el)
    img_p = Paragraph(img_el, anchor._parent)
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(str(image), width=Inches(5.8))

    cap_el = OxmlElement("w:p")
    img_el.addnext(cap_el)
    cap_p = Paragraph(cap_el, anchor._parent)
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap_p.add_run(caption)
    _set_run_font(r, size=10.5)
    return cap_p


def main() -> int:
    doc = docx.Document(str(SRC))
    report_cell = doc.tables[33].rows[2].cells[0]
    paragraphs = report_cell.paragraphs

    inserted = 0
    last_for_section: dict[str, Paragraph] = {}
    for prefix, img_name, caption in PLAN:
        anchor = last_for_section.get(prefix)
        if anchor is None:
            anchor = next((p for p in paragraphs if p.text.startswith(prefix)), None)
        if anchor is None:
            print("WARN: section not found:", prefix)
            continue
        cap_p = _insert_image(report_cell, anchor, IMG_DIR / img_name, caption)
        last_for_section[prefix] = cap_p
        inserted += 1

    doc.save(str(DST))
    print("inserted images:", inserted)
    print("saved:", DST)

    # validate
    chk = docx.Document(str(DST))
    print("inline_shapes:", len(chk.inline_shapes))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
